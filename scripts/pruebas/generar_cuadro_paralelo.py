#!/usr/bin/env python3
"""Genera únicamente el cuadro PAR en Word, en formato vertical y sin colores."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# Una fila por módulo visible/submódulo funcional del sistema.
# Los endpoints se han contrastado con los controladores reales.
# N/A indica que no hay una ruta API equivalente para esa operación.
ROWS = [
    ["N/A", "Inicio de sesión", "IniciarSesion/mostrarIniciarSesion", "IniciarSesion/iniciarSesionMovilApk (APK)", "Fuera del alcance móvil."],
    ["N/A", "Recuperación de contraseña", "RecuperarContr/mostrarRecuperarContr", "RecuperarContr/*Apk (APK)", "Fuera del alcance móvil."],
    ["N/A", "Inicio / Dashboard", "Inicio/inicio", "N/A", "N/A"],
    ["PAR-01", "Pacientes: búsqueda por cédula", "Pacientes/getPacientesAjax", "Api/apiBuscarPaciente?q=<dato>", "Mismo paciente: cédula, nombre y apellido."],
    ["N/A", "Clientes: listado y papelera", "Clientes/Clientes; Clientes/papelera", "N/A", "N/A"],
    ["N/A", "Patologías: catálogo y papelera", "Patologias/patologias; Patologias/papeleraPatologias", "N/A", "N/A"],
    ["PAR-05", "Patologías por paciente", "Control/mostrarControlPacientesJS/<cédula>", "Api/apiPatologiasPaciente/<id>", "Mismas patologías y fechas."],
    ["N/A", "Facturación: general, cita y hospitalización", "Factura/factura; facturaCita; facturarHospitalizacion", "N/A", "N/A"],
    ["N/A", "Citas: pendientes, hoy y realizadas", "Citas/citas; citasHoy; citasRealizadas", "Citas/citasHoyCompletasApk (APK)", "Fuera del alcance móvil."],
    ["PAR-03", "Citas: doctores por especialidad", "Citas/mostrarDoctoresCita/<hash>", "Api/apiDoctoresPorEspecialidad/<id>", "Mismos doctores: nombre y apellido."],
    ["PAR-02", "Servicios: categorías activas", "Servicios/categoriasAjax", "Api/apiEspecialidades", "Mismos nombres y cantidad."],
    ["N/A", "Directorio médico: listado y papelera", "Doctores/doctores; Doctores/papelera", "N/A", "N/A"],
    ["PAR-04", "Control médico: controles del paciente", "Control/mostrarControlPacientesJS/<cédula>", "Api/apiControlesPaciente/<id>", "Mismos controles comparables."],
    ["N/A", "Hospitalización: en curso y realizadas", "Hospitalizacion/hospitalizacion; hospitalizacionesRealizadas", "Hospitalizacion/hospitalizacionApk (APK)", "Fuera del alcance móvil."],
    ["N/A", "Insumos: inventario, vencidos y papelera", "Insumos/insumos; InsumosVencidos; papelera", "N/A", "N/A"],
    ["N/A", "Entradas: activas y desactivadas", "Entrada/entrada; Entrada/papelera", "N/A", "N/A"],
    ["N/A", "Proveedores: listado y papelera", "Proveedores/proveedores; Proveedores/papelera", "N/A", "N/A"],
    ["N/A", "Usuarios: listado", "Usuarios/usuarios", "N/A", "N/A"],
    ["N/A", "Administradores", "Usuarios/administradores", "N/A", "N/A"],
    ["N/A", "Roles y permisos", "Roles/mostrar", "N/A", "N/A"],
    ["N/A", "Lista negra de usuarios", "Usuarios/listaNegra", "N/A", "N/A"],
    ["N/A", "Bitácora: usuario y administrador", "Bitacora/bitacoraUsuario; Bitacora/bitacora", "N/A", "N/A"],
    ["N/A", "Reportes y exportaciones", "Reportes/reportes", "N/A", "N/A"],
    ["N/A", "Estadísticas", "Estadisticas/estadisticas", "N/A", "N/A"],
    ["N/A", "Mantenimiento y respaldos", "Mantenimiento/mantenimiento", "N/A", "N/A"],
    ["N/A", "Perfil", "Perfil/perfil", "Perfil/perfilApk; editarPerfilApk (APK)", "Fuera del alcance móvil."],
    ["N/A", "Asistente IA", "Modal del asistente en Inicio", "Chat/preguntar (mismo endpoint que usa la interfaz)", "No hay una segunda vía independiente."],
]


def build(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)
    heading.add_run("4.2.12 Datos de las plantillas de pruebas Paralelas:").font.size = Pt(11)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [0.45, 1.25, 1.75, 1.9, 2.15]
    headers = ["ID", "Caso propuesto", "Vía web", "Vía API", "Criterio de aprobación"]

    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(label)
        run.font.name = "Arial"
        run.font.size = Pt(9)

    for item in ROWS:
        cells = table.add_row().cells
        for index, value in enumerate(item):
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(8)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    default_output = (
        Path.home()
        / "Documents"
        / "universidad"
        / "proyecto"
        / "Clinica"
        / "Cuadro_Paralela_Todos_Modulos_v2.docx"
    )
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--salida", type=Path, default=default_output)
    args = parser.parse_args()
    build(args.salida)
    print(f"Generado: {args.salida.resolve()}")
    print(f"Módulos/áreas en el cuadro: {len(ROWS)}")
    print(f"Comparaciones PAR ejecutables con las rutas /Api actuales: {sum(r[0] != 'N/A' for r in ROWS)}")


if __name__ == "__main__":
    main()
